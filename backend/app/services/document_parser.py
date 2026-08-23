"""将首期支持的运营资料解析为可检索的纯文本片段。"""

from dataclasses import dataclass
from io import BytesIO


SUPPORTED_FILE_TYPES = {".pdf", ".docx", ".md", ".txt"}


@dataclass(frozen=True)
class ParsedChunk:
    """一段用于检索的文本及其在原始资料中的位置。"""

    content: str
    page_number: int | None = None
    heading: str | None = None


@dataclass(frozen=True)
class ParsedDocument:
    """解析后的完整正文和按语义边界切分的文本片段。"""

    content: str
    chunks: list[ParsedChunk]


def parse_document(filename: str, raw_content: bytes) -> ParsedDocument:
    """根据扩展名提取正文；不支持的格式必须在写库前被拒绝。"""
    suffix = _suffix(filename)
    if suffix not in SUPPORTED_FILE_TYPES:
        raise ValueError("仅支持 PDF、DOCX、Markdown 和 TXT 文件")
    if suffix == ".pdf":
        return _parse_pdf(raw_content)
    if suffix == ".docx":
        return _parse_docx(raw_content)
    return _parse_text(raw_content)


def _parse_pdf(raw_content: bytes) -> ParsedDocument:
    import fitz

    document = fitz.open(stream=raw_content, filetype="pdf")
    page_chunks: list[ParsedChunk] = []
    pages: list[str] = []
    try:
        for page_number, page in enumerate(document, start=1):
            page_text = page.get_text("text").strip()
            if not page_text:
                continue
            pages.append(page_text)
            page_chunks.extend(
                ParsedChunk(content=chunk, page_number=page_number)
                for chunk in _split_text(page_text)
            )
    finally:
        document.close()
    return _finish("\n\n".join(pages), page_chunks)


def _parse_docx(raw_content: bytes) -> ParsedDocument:
    from docx import Document

    document = Document(BytesIO(raw_content))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    return _finish("\n\n".join(paragraphs), [ParsedChunk(content=chunk) for chunk in _split_text("\n\n".join(paragraphs))])


def _parse_text(raw_content: bytes) -> ParsedDocument:
    try:
        text = raw_content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ValueError("文本文件必须使用 UTF-8 编码") from error
    text = text.strip()
    return _finish(text, [ParsedChunk(content=chunk) for chunk in _split_text(text)])


def _finish(content: str, chunks: list[ParsedChunk]) -> ParsedDocument:
    if not content or not chunks:
        raise ValueError("文件中未提取到可检索的文本内容")
    return ParsedDocument(content=content, chunks=chunks)


def _split_text(text: str, chunk_size: int = 800, overlap: int = 120) -> list[str]:
    """优先在段落或句末切分，保留少量重叠以避免上下文被截断。"""
    normalized = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    chunks: list[str] = []
    start = 0
    while start < len(normalized):
        end = min(start + chunk_size, len(normalized))
        if end < len(normalized):
            boundary = max(normalized.rfind("\n", start, end), normalized.rfind("。", start, end))
            if boundary > start + chunk_size // 2:
                end = boundary + 1
        chunk = normalized[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(normalized):
            break
        start = max(end - overlap, start + 1)
    return chunks


def _suffix(filename: str) -> str:
    return f".{filename.rsplit('.', 1)[-1].lower()}" if "." in filename else ""
