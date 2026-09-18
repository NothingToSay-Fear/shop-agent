"""将运营资料解析为保留章节结构的可检索文本片段。"""

from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO
from typing import Literal


SUPPORTED_FILE_TYPES = {".pdf", ".docx", ".md", ".txt"}
CHUNK_MIN_SIZE = 300
CHUNK_TARGET_SIZE = 650
CHUNK_MAX_SIZE = 900
HARD_SPLIT_OVERLAP = 100
ChunkContentType = Literal["paragraph", "list", "table", "mixed"]
_LIST_PATTERN = re.compile(r"^\s*(?:[-*+]\s+|\d+[.)、]\s+|[一二三四五六七八九十]+[、.])")
_TEXT_HEADING_PATTERN = re.compile(r"^\s*(第[一二三四五六七八九十\d]+[章节]|[一二三四五六七八九十]+、)\s*(.+)$")
_SENTENCE_BOUNDARY = re.compile(r"[。！？；;]\s*")


@dataclass(frozen=True)
class ParsedBlock:
    """一个尚未合并的结构块，例如段落、连续列表或表格。"""

    content: str
    heading_path: tuple[str, ...] = ()
    content_type: ChunkContentType = "paragraph"
    page_start: int | None = None
    page_end: int | None = None


@dataclass(frozen=True)
class ParsedChunk:
    """一段用于检索的文本及其在原始资料中的结构定位。"""

    content: str
    heading: str | None = None
    heading_path: str | None = None
    content_type: ChunkContentType = "paragraph"
    page_start: int | None = None
    page_end: int | None = None


@dataclass(frozen=True)
class ParsedDocument:
    """解析后的全文、结构块与初步按章节组装的检索片段。"""

    content: str
    blocks: list[ParsedBlock]
    chunks: list[ParsedChunk]


def parse_document(filename: str, raw_content: bytes) -> ParsedDocument:
    """根据扩展名提取结构块；不支持的格式必须在写库前被拒绝。"""
    suffix = _suffix(filename)
    if suffix not in SUPPORTED_FILE_TYPES:
        raise ValueError("仅支持 PDF、DOCX、Markdown 和 TXT 文件")
    if suffix == ".pdf":
        return _parse_pdf(raw_content)
    if suffix == ".docx":
        return _parse_docx(raw_content)
    return _parse_text(raw_content, markdown=suffix == ".md")


async def apply_semantic_boundaries(
    parsed: ParsedDocument,
    semantic_similarity_threshold: float,
    embedder,
    batch_size: int = 32,
) -> ParsedDocument:
    """用本地段落向量辅助切分；模型不可用时保留确定性的结构化结果。"""
    if len(parsed.blocks) < 2:
        return parsed
    embeddings: list[list[float]] = []
    texts = [block.content for block in parsed.blocks]
    for offset in range(0, len(texts), max(batch_size, 1)):
        batch = await embedder(texts[offset : offset + max(batch_size, 1)])
        if batch is None or len(batch) != len(texts[offset : offset + max(batch_size, 1)]):
            return parsed
        embeddings.extend(batch)
    similarities: dict[int, float] = {}
    for index, (left, right) in enumerate(zip(embeddings, embeddings[1:]), start=1):
        if len(left) != len(right) or not left:
            continue
        similarities[index] = sum(a * b for a, b in zip(left, right, strict=True))
    return ParsedDocument(
        content=parsed.content,
        blocks=parsed.blocks,
        chunks=_build_chunks(parsed.blocks, similarities, semantic_similarity_threshold),
    )


def _parse_pdf(raw_content: bytes) -> ParsedDocument:
    import fitz

    document = fitz.open(stream=raw_content, filetype="pdf")
    blocks: list[ParsedBlock] = []
    pages: list[str] = []
    try:
        for page_number, page in enumerate(document, start=1):
            page_text = page.get_text("text").strip()
            if not page_text:
                continue
            pages.append(page_text)
            blocks.extend(_plain_text_blocks(page_text, page_start=page_number, page_end=page_number))
    finally:
        document.close()
    return _finish("\n\n".join(pages), blocks)


def _parse_docx(raw_content: bytes) -> ParsedDocument:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(BytesIO(raw_content))
    blocks: list[ParsedBlock] = []
    heading_stack: list[str] = []
    full_text: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            level = _docx_heading_level(paragraph.style.name if paragraph.style else "")
            if level is not None:
                _update_heading_stack(heading_stack, level, text)
                full_text.append(text)
                continue
            full_text.append(text)
            blocks.extend(_plain_text_blocks(text, tuple(heading_stack)))
        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            table_text = _docx_table_text(table)
            if table_text:
                full_text.append(table_text)
                blocks.append(
                    ParsedBlock(
                        content=table_text,
                        heading_path=tuple(heading_stack),
                        content_type="table",
                    )
                )
    return _finish("\n\n".join(full_text), blocks)


def _parse_text(raw_content: bytes, markdown: bool) -> ParsedDocument:
    try:
        text = raw_content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ValueError("文本文件必须使用 UTF-8 编码") from error
    text = text.strip()
    blocks = _markdown_blocks(text) if markdown else _plain_text_blocks(text, detect_text_headings=True)
    return _finish(text, blocks)


def _markdown_blocks(text: str) -> list[ParsedBlock]:
    """识别 Markdown 标题、连续列表和表格，普通连续行合并为段落。"""
    blocks: list[ParsedBlock] = []
    heading_stack: list[str] = []
    pending: list[str] = []
    pending_type: ChunkContentType = "paragraph"

    def flush() -> None:
        nonlocal pending
        content = "\n".join(pending).strip()
        if content:
            blocks.append(ParsedBlock(content, tuple(heading_stack), pending_type))
        pending = []

    for line in text.splitlines():
        stripped = line.strip()
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", stripped)
        if heading:
            flush()
            _update_heading_stack(heading_stack, len(heading.group(1)), heading.group(2).strip())
            continue
        if not stripped:
            flush()
            continue
        line_type: ChunkContentType = "table" if _is_markdown_table_line(stripped) else (
            "list" if _LIST_PATTERN.match(stripped) else "paragraph"
        )
        if pending and line_type != pending_type:
            flush()
        pending_type = line_type
        pending.append(stripped)
    flush()
    return blocks


def _plain_text_blocks(
    text: str,
    heading_path: tuple[str, ...] = (),
    page_start: int | None = None,
    page_end: int | None = None,
    detect_text_headings: bool = False,
) -> list[ParsedBlock]:
    """按空行与连续列表组织纯文本；TXT 可识别简单章节编号。"""
    blocks: list[ParsedBlock] = []
    current_heading = list(heading_path)
    pending: list[str] = []
    pending_type: ChunkContentType = "paragraph"

    def flush() -> None:
        nonlocal pending
        content = "\n".join(pending).strip()
        if content:
            blocks.append(
                ParsedBlock(content, tuple(current_heading), pending_type, page_start, page_end)
            )
        pending = []

    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            flush()
            continue
        heading = _TEXT_HEADING_PATTERN.match(stripped) if detect_text_headings else None
        if heading:
            flush()
            _update_heading_stack(current_heading, 1, stripped)
            continue
        line_type: ChunkContentType = "list" if _LIST_PATTERN.match(stripped) else "paragraph"
        if pending and line_type != pending_type:
            flush()
        pending_type = line_type
        pending.append(stripped)
    flush()
    return blocks


def _build_chunks(
    blocks: list[ParsedBlock],
    semantic_similarities: dict[int, float] | None = None,
    semantic_similarity_threshold: float | None = None,
) -> list[ParsedChunk]:
    """在同一章节内合并完整结构块，仅对超长块进行句级硬切。"""
    chunks: list[ParsedChunk] = []
    current: list[ParsedBlock] = []

    def flush() -> None:
        nonlocal current
        if current:
            chunks.append(_make_chunk(current))
        current = []

    for index, block in enumerate(blocks):
        block_parts = _split_oversized_block(block)
        for part_index, part in enumerate(block_parts):
            if not current:
                current.append(part)
                continue
            previous = current[-1]
            boundary_similarity = (semantic_similarities or {}).get(index) if part_index == 0 else None
            must_split = (
                previous.heading_path != part.heading_path
                or previous.page_end != part.page_start
                or previous.content_type == "table"
                or part.content_type == "table"
            )
            current_size = _blocks_length(current)
            candidate_size = current_size + 2 + len(part.content)
            semantic_break = (
                semantic_similarity_threshold is not None
                and boundary_similarity is not None
                and current_size >= CHUNK_MIN_SIZE
                and boundary_similarity < semantic_similarity_threshold
            )
            if must_split or semantic_break or (
                current_size >= CHUNK_TARGET_SIZE and candidate_size > CHUNK_TARGET_SIZE
            ) or candidate_size > CHUNK_MAX_SIZE:
                flush()
            current.append(part)
    flush()
    return chunks


def _split_oversized_block(block: ParsedBlock) -> list[ParsedBlock]:
    """超长段落按句子切分；表格按行组切分并在后续片段重复表头。"""
    if len(block.content) <= CHUNK_MAX_SIZE:
        return [block]
    if block.content_type == "table":
        return _split_table_block(block)
    sentences = _SENTENCE_BOUNDARY.split(block.content)
    sentences = [sentence.strip() for sentence in sentences if sentence.strip()]
    if len(sentences) <= 1:
        return _hard_split(block)
    parts: list[ParsedBlock] = []
    current = ""
    for sentence in sentences:
        candidate = f"{current}。{sentence}" if current else sentence
        if current and len(candidate) > CHUNK_MAX_SIZE:
            parts.append(_replace_block_content(block, current))
            current = sentence
        else:
            current = candidate
    if current:
        parts.append(_replace_block_content(block, current))
    return parts or _hard_split(block)


def _split_table_block(block: ParsedBlock) -> list[ParsedBlock]:
    lines = [line for line in block.content.splitlines() if line.strip()]
    if len(lines) <= 1:
        return _hard_split(block)
    header = lines[0]
    parts: list[ParsedBlock] = []
    current = header
    for line in lines[1:]:
        candidate = f"{current}\n{line}"
        if len(candidate) > CHUNK_MAX_SIZE and current != header:
            parts.append(_replace_block_content(block, current))
            current = f"{header}\n{line}"
        else:
            current = candidate
    if current:
        parts.append(_replace_block_content(block, current))
    return parts


def _hard_split(block: ParsedBlock) -> list[ParsedBlock]:
    parts: list[ParsedBlock] = []
    start = 0
    while start < len(block.content):
        end = min(start + CHUNK_MAX_SIZE, len(block.content))
        if end < len(block.content):
            boundary = max(block.content.rfind("\n", start, end), block.content.rfind("。", start, end))
            if boundary > start + CHUNK_MIN_SIZE:
                end = boundary + 1
        parts.append(_replace_block_content(block, block.content[start:end].strip()))
        if end >= len(block.content):
            break
        start = max(end - HARD_SPLIT_OVERLAP, start + 1)
    return parts


def _make_chunk(blocks: list[ParsedBlock]) -> ParsedChunk:
    heading_path = blocks[0].heading_path
    content_types = {block.content_type for block in blocks}
    return ParsedChunk(
        content="\n\n".join(block.content for block in blocks),
        heading=heading_path[-1] if heading_path else None,
        heading_path=" > ".join(heading_path) or None,
        content_type=content_types.pop() if len(content_types) == 1 else "mixed",
        page_start=blocks[0].page_start,
        page_end=blocks[-1].page_end,
    )


def _replace_block_content(block: ParsedBlock, content: str) -> ParsedBlock:
    return ParsedBlock(
        content=content,
        heading_path=block.heading_path,
        content_type=block.content_type,
        page_start=block.page_start,
        page_end=block.page_end,
    )


def _blocks_length(blocks: list[ParsedBlock]) -> int:
    return sum(len(block.content) for block in blocks) + max(len(blocks) - 1, 0) * 2


def _update_heading_stack(stack: list[str], level: int, heading: str) -> None:
    while len(stack) >= level:
        stack.pop()
    stack.append(heading)


def _docx_heading_level(style_name: str) -> int | None:
    match = re.search(r"(?:heading|标题)\s*(\d+)", style_name, re.IGNORECASE)
    return int(match.group(1)) if match else None


def _docx_table_text(table) -> str:
    rows = []
    for row in table.rows:
        values = [" ".join(cell.text.split()) for cell in row.cells]
        if any(values):
            rows.append(" | ".join(values))
    return "\n".join(rows)


def _is_markdown_table_line(line: str) -> bool:
    return "|" in line and (line.startswith("|") or line.endswith("|"))


def _finish(content: str, blocks: list[ParsedBlock]) -> ParsedDocument:
    if not content or not blocks:
        raise ValueError("文件中未提取到可检索的文本内容")
    return ParsedDocument(content=content, blocks=blocks, chunks=_build_chunks(blocks))


def _suffix(filename: str) -> str:
    return f".{filename.rsplit('.', 1)[-1].lower()}" if "." in filename else ""
