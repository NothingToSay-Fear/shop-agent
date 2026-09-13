import pytest

from app.services.document_parser import apply_semantic_boundaries, parse_document


def test_text_document_is_split_into_searchable_chunks() -> None:
    """UTF-8 文本资料应保留全文，并切分为至少一个可检索片段。"""
    parsed = parse_document("活动规则.md", "报名时间为 8 月 1 日至 8 月 7 日。\n商品需要满足库存要求。".encode())

    assert "报名时间" in parsed.content
    assert parsed.chunks[0].page_start is None
    assert "库存要求" in parsed.chunks[0].content


def test_pdf_document_keeps_page_number() -> None:
    """PDF 资料片段必须保留页码，以便问答结果给出可复核引用。"""
    import fitz

    source = fitz.open()
    page = source.new_page()
    # PyMuPDF 的默认测试字体不含中文；真实中文 PDF 会携带自身的字体与文本映射。
    page.insert_text((72, 72), "Registration requires store verification.")
    raw_content = source.tobytes()
    source.close()

    parsed = parse_document("规则.pdf", raw_content)

    assert parsed.chunks[0].page_start == 1
    assert parsed.chunks[0].page_end == 1
    assert "Registration" in parsed.chunks[0].content


def test_docx_document_extracts_paragraphs() -> None:
    """DOCX 资料应从段落中提取可检索文本。"""
    from docx import Document
    from io import BytesIO

    document = Document()
    document.add_paragraph("优惠券仅限活动期间领取。")
    buffer = BytesIO()
    document.save(buffer)

    parsed = parse_document("规则.docx", buffer.getvalue())

    assert "优惠券" in parsed.content


def test_markdown_chunk_keeps_heading_path_and_list_together() -> None:
    """Markdown 的标题路径和连续规则列表应保留到同一个可检索片段。"""
    parsed = parse_document(
        "规则.md",
        (
            "# 2026 年 618 活动规则\n\n"
            "## 跨店满减\n\n"
            "- 满 300 减 50\n- 满 600 减 100\n- 满 900 减 150\n"
        ).encode(),
    )

    assert parsed.chunks[0].heading_path == "2026 年 618 活动规则 > 跨店满减"
    assert parsed.chunks[0].content_type == "list"
    assert "满 300 减 50" in parsed.chunks[0].content
    assert "满 900 减 150" in parsed.chunks[0].content


def test_docx_chunk_keeps_heading_and_table() -> None:
    """DOCX 标题样式和表格应生成带章节定位的独立结构块。"""
    from docx import Document
    from io import BytesIO

    document = Document()
    document.add_heading("优惠玩法", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "门槛"
    table.cell(0, 1).text = "优惠"
    table.cell(1, 0).text = "满300"
    table.cell(1, 1).text = "减50"
    buffer = BytesIO()
    document.save(buffer)

    parsed = parse_document("规则.docx", buffer.getvalue())

    assert parsed.chunks[0].heading_path == "优惠玩法"
    assert parsed.chunks[0].content_type == "table"
    assert "满300 | 减50" in parsed.chunks[0].content


@pytest.mark.asyncio
async def test_semantic_boundary_splits_low_similarity_paragraphs() -> None:
    """达到最小长度后，主题明显跳变的相邻段落应被本地向量切开。"""
    parsed = parse_document(
        "资料.txt",
        (("活动报名条件说明。" * 40) + "\n\n" + ("售后退款处理流程。" * 40)).encode(),
    )

    async def fake_embedder(_texts: list[str]) -> list[list[float]]:
        return [[1.0, 0.0], [0.0, 1.0]]

    refined = await apply_semantic_boundaries(parsed, 0.55, fake_embedder)

    assert len(refined.chunks) == 2


def test_unsupported_document_type_is_rejected() -> None:
    """未实现解析器的文件不能被错误地纳入知识库。"""
    try:
        parse_document("活动规则.xlsx", b"not an excel file")
    except ValueError as error:
        assert "仅支持" in str(error)
    else:
        raise AssertionError("应拒绝不支持的文件格式")
